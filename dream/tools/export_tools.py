"""Exporters: the artifact leaves the frame.

- `super_inline_html`: one self-contained file, assets inlined (dream.gui.bundle).
- `gen_pptx`: the deck in the hidden frame becomes a .pptx — `screenshots` mode is
  one full-bleed picture per slide; `editable` mode reads each slide's DOM and
  emits native text boxes, filled rectangles, and pictures at their positions,
  with a screenshot behind them only if asked. Speaker notes come from the page's
  `#speaker-notes` JSON. Validation flags say what a bad capture looks like.
"""

from __future__ import annotations

import base64
import io
import json
import re
import time
from pathlib import Path
from typing import Any

from claude_agent_sdk import tool

from .. import config
from ..gui.preview import get_preview
from . import mirror
from .context import ctx, err, in_thread, ok

_EMU_PER_PX = 9525  # 96 dpi


def _resolve(raw: str) -> Path:
    return (ctx().workspace / Path(raw).expanduser()).resolve()


def _inside(p: Path) -> bool:
    return p.is_relative_to(ctx().workspace.resolve())


# --- super_inline_html --------------------------------------------------------------------


@tool(
    "super_inline_html",
    "Bundle an HTML file and all its referenced local assets (scripts, stylesheets, images, "
    "fonts) into a single self-contained HTML file that works offline. The input MUST "
    "contain a <template id=\"__bundler_thumbnail\"> with a simple colorful SVG splash "
    "(an icon or 1–2 letters, generous padding) — shown while the bundle unpacks and as "
    "the no-JS fallback. Remote URLs and files outside the workspace are left as they "
    "are and listed.",
    {"type": "object", "properties": {"input_path": {"type": "string"}, "output_path": {"type": "string"}},
     "required": ["input_path", "output_path"]},
)
async def super_inline_html(args: dict[str, Any]) -> dict[str, Any]:
    src, dst = args.get("input_path"), args.get("output_path")
    if not src or not dst:
        return err("super_inline_html needs input_path and output_path.")
    s, d = _resolve(str(src)), _resolve(str(dst))
    if not s.is_file():
        return err(f"No file at {s}")
    if not _inside(d):
        return err("output_path must be inside the workspace.")
    from ..gui.bundle import bundle

    try:
        html, report = await in_thread(bundle, s, ctx().workspace.resolve())
    except ValueError as e:
        return err(str(e))
    except Exception as e:
        return err(f"Bundling failed: {type(e).__name__}: {e}")
    d.parent.mkdir(parents=True, exist_ok=True)
    d.write_text(html, encoding="utf-8")
    mirror.file_written(d)   # a bundle written over the shown page reloads it (DREAM-104)
    inlined = sum(r.startswith("inlined") for r in report)
    left = [r for r in report if r.startswith("left alone")]
    return ok(f"Wrote {d} ({len(html):,} chars): {inlined} asset(s) inlined."
              + ("\nLeft as-is:\n" + "\n".join(left) if left else ""))


# --- gen_pptx ---------------------------------------------------------------------------------


_SLIDE_DOM_JS = r"""
(sel) => {
  const root = document.querySelector(sel);
  if (!root) return null;
  const rb = root.getBoundingClientRect();
  const out = { box: {x: rb.x, y: rb.y, w: rb.width, h: rb.height}, items: [] };
  const seen = new Set();
  const walk = (el) => {
    if (!(el instanceof Element)) return;
    const cs = getComputedStyle(el);
    if (cs.display === 'none' || cs.visibility === 'hidden') return;
    const b = el.getBoundingClientRect();
    if (b.width < 1 || b.height < 1) return;
    const rel = {x: b.x - rb.x, y: b.y - rb.y, w: b.width, h: b.height};
    const bg = cs.backgroundColor;
    // the slide root's own background is the slide background: a full rect
    if (bg && !/rgba\(\s*0,\s*0,\s*0,\s*0\)/.test(bg) && bg !== 'transparent') {
      out.items.push({kind: 'rect', ...rel, fill: bg, radius: parseFloat(cs.borderTopLeftRadius) || 0});
    }
    if (el.tagName === 'IMG' && el.currentSrc) {
      // Prefer pixels: a same-origin (file:, data:) image drawn to a canvas comes
      // back as a data URL; anything else keeps its src for the Python side.
      let src = el.currentSrc;
      try {
        const c = document.createElement('canvas');
        c.width = el.naturalWidth || Math.round(b.width); c.height = el.naturalHeight || Math.round(b.height);
        c.getContext('2d').drawImage(el, 0, 0, c.width, c.height);
        src = c.toDataURL('image/png');
      } catch (e) {}
      out.items.push({kind: 'image', ...rel, src, origin: el.currentSrc});
    }
    if (el.tagName === 'SVG' || el.tagName === 'svg') {
      out.items.push({kind: 'svg', ...rel, xml: new XMLSerializer().serializeToString(el)});
      return;
    }
    // direct text of this element (not descendants), as one box
    let text = '';
    for (const n of el.childNodes) if (n.nodeType === 3) text += n.textContent;
    text = text.replace(/\s+/g, ' ').trim();
    if (text) {
      const range = document.createRange(); range.selectNodeContents(el);
      const tb = range.getBoundingClientRect();
      out.items.push({kind: 'text', x: tb.x - rb.x, y: tb.y - rb.y, w: tb.width, h: tb.height, text,
        size: parseFloat(cs.fontSize), weight: cs.fontWeight, italic: cs.fontStyle === 'italic',
        color: cs.color, family: cs.fontFamily, align: cs.textAlign});
    }
    for (const c of el.children) walk(c);
  };
  walk(root);
  return out;
}
"""


def _rgb(css: str) -> tuple[int, int, int] | None:
    m = re.match(r"rgba?\(\s*(\d+),\s*(\d+),\s*(\d+)", css or "")
    return (int(m.group(1)), int(m.group(2)), int(m.group(3))) if m else None


def _file_of(url: str) -> Path:
    from urllib.parse import unquote, urlparse

    return Path(unquote(urlparse(url).path)).resolve()


def _image_bytes(src: str, workspace: Path, origin: str = "") -> bytes | None:
    """A data: URL's bytes, or a file: URL's bytes if the file is inside the
    workspace. Anything else (remote, outside) is skipped, never fetched. `origin`
    is where the pixels came from when `src` is a canvas copy — a file outside
    the workspace stays outside even after the page drew it."""
    if origin.startswith("file://") and not _file_of(origin).is_relative_to(workspace):
        return None
    if src.startswith("data:image/"):
        try:
            return base64.b64decode(src.split(",", 1)[1])
        except Exception:
            return None
    if src.startswith("file://"):
        p = _file_of(src)
        if p.is_file() and p.is_relative_to(workspace):
            return p.read_bytes()
    return None


def _build_pptx(width: int, height: int, slides: list[dict[str, Any]], notes: list[str],
                mode: str, workspace: Path) -> tuple[bytes, list[str]]:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Emu, Pt

    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(width * _EMU_PER_PX), Emu(height * _EMU_PER_PX)
    blank = prs.slide_layouts[6]
    flags: list[str] = []
    prev_png: bytes | None = None
    for i, s in enumerate(slides):
        slide = prs.slides.add_slide(blank)
        png = s.get("png")
        if png is not None and prev_png is not None and png == prev_png:
            flags.append(f"duplicate_adjacent: slide {i + 1} captured identical to slide {i} — "
                         f"did showJs navigate?")
        prev_png = png
        box = s.get("box") or {}
        if box and (abs(box.get("w", width) - width) > 2 or abs(box.get("h", height) - height) > 2):
            flags.append(f"slide_size_mismatch: slide {i + 1} selector box is "
                         f"{box.get('w'):.0f}×{box.get('h'):.0f}, not {width}×{height} — wrong "
                         f"selector or resetTransformSelector?")
        if mode == "screenshots" or s.get("background"):
            if png is not None:
                slide.shapes.add_picture(io.BytesIO(png), 0, 0, prs.slide_width, prs.slide_height)
        if mode == "editable":
            for it in s.get("items") or []:
                x, y, w, h = (Emu(int(it[k] * _EMU_PER_PX)) for k in ("x", "y", "w", "h"))
                if it["kind"] == "rect":
                    rgb = _rgb(it.get("fill", ""))
                    if rgb is None:
                        continue
                    shp = slide.shapes.add_shape(
                        MSO_SHAPE.ROUNDED_RECTANGLE if it.get("radius") else MSO_SHAPE.RECTANGLE, x, y, w, h)
                    shp.fill.solid(); shp.fill.fore_color.rgb = RGBColor(*rgb)
                    shp.line.fill.background()
                elif it["kind"] == "image":
                    data = _image_bytes(str(it.get("src", "")), workspace, str(it.get("origin", "")))
                    if data is None:
                        flags.append(f"image_skipped: slide {i + 1} had an image that could not be read "
                                     f"({str(it.get('src', ''))[:60]})")
                        continue
                    try:
                        slide.shapes.add_picture(io.BytesIO(data), x, y, w, h)
                    except Exception:
                        flags.append(f"image_skipped: slide {i + 1} had an image that could not be decoded")
                elif it["kind"] == "text":
                    tb = slide.shapes.add_textbox(x, y, w, h)
                    tf = tb.text_frame; tf.word_wrap = True
                    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
                    run = tf.paragraphs[0].add_run(); run.text = it["text"]
                    run.font.size = Pt(max(1, it.get("size", 16) * 0.75))
                    run.font.bold = str(it.get("weight", "")).isdigit() and int(it["weight"]) >= 600 or it.get("weight") == "bold"
                    run.font.italic = bool(it.get("italic"))
                    fam = str(it.get("family", "")).split(",")[0].strip().strip("'\"")
                    if fam:
                        run.font.name = fam
                    rgb = _rgb(it.get("color", ""))
                    if rgb:
                        run.font.color.rgb = RGBColor(*rgb)
        if i < len(notes) and notes[i]:
            slide.notes_slide.notes_text_frame.text = str(notes[i])
    if not any(notes):
        flags.append("no_speaker_notes: no #speaker-notes script found (fine if the deck has none)")
    buf = io.BytesIO(); prs.save(buf)
    return buf.getvalue(), flags


@tool(
    "gen_pptx",
    "Export the deck loaded in the hidden frame to a .pptx. Give the slide size in CSS px "
    "(1920×1080 for the deck starter) and one entry per slide: a `selector` for the "
    "slide's root element and optional `showJs` to navigate to it first (e.g. "
    "\"goToSlide(2)\"; sync expression, the per-slide `delay` covers transitions). "
    "`resetTransformSelector` clears the scaling on a scaled deck — pass \"deck-stage\" for "
    "the starter; it drops its scale when `noscale` is set. Modes: `screenshots` (one "
    "full-bleed PNG per slide, pixel-perfect, not editable) or `editable` (native text "
    "boxes, filled shapes, and pictures at their positions; `background` adds the "
    "screenshot behind them). Speaker notes are read from `<script type=\"application/json\" "
    "id=\"speaker-notes\">`. Returns validation flags: read each one and decide whether it "
    "is expected for THIS deck — duplicate_adjacent means showJs probably did not "
    "navigate; slide_size_mismatch means the selector or resetTransformSelector is wrong; "
    "no_speaker_notes is fine for a deck without notes.",
    {
        "type": "object",
        "properties": {
            "path": {"type": "string", "description": "The deck's HTML file (loaded if not already)."},
            "width": {"type": "number"}, "height": {"type": "number"},
            "slides": {"type": "array", "items": {"type": "object", "properties": {
                "selector": {"type": "string"}, "showJs": {"type": "string"}, "delay": {"type": "number"}},
                "required": ["selector"]}},
            "mode": {"type": "string", "enum": ["screenshots", "editable"]},
            "background": {"type": "boolean", "description": "editable mode: put the screenshot behind the shapes."},
            "resetTransformSelector": {"type": "string"},
            "hideSelectors": {"type": "array", "items": {"type": "string"}},
            "filename": {"type": "string", "description": "Output name without extension (default deck), in the workspace."},
        },
        "required": ["path", "width", "height", "slides"],
    },
)
async def gen_pptx(args: dict[str, Any]) -> dict[str, Any]:
    raw = args.get("path")
    if not raw:
        return err("gen_pptx needs 'path' (the deck's HTML file).")
    p = _resolve(str(raw))
    if not p.is_file():
        return err(f"No file at {p}")
    try:
        width, height = int(args["width"]), int(args["height"])
    except (KeyError, TypeError, ValueError):
        return err("gen_pptx needs numeric width and height.")
    slides = args.get("slides")
    if not isinstance(slides, list) or not slides or not all(isinstance(s, dict) and s.get("selector") for s in slides):
        return err("gen_pptx needs a non-empty 'slides' list; each entry needs a selector.")
    mode = str(args.get("mode") or "screenshots")
    if mode not in ("screenshots", "editable"):
        return err("mode must be 'screenshots' or 'editable'.")
    name = re.sub(r"[^A-Za-z0-9._-]+", "_", str(args.get("filename") or "deck")).strip("._") or "deck"
    out = ctx().workspace.resolve() / f"{name}.pptx"

    pv = get_preview()
    try:
        if pv.loaded != p.resolve():
            await pv.load(p)
        rts = args.get("resetTransformSelector")
        if rts:
            # `noscale` is the contract the deck starter honours; the inline style
            # covers a hand-rolled deck that scales with a transform of its own.
            await pv.eval(
                f"(() => {{ const el = document.querySelector({json.dumps(str(rts))}); if (!el) return false; "
                f"el.setAttribute('noscale', ''); el.style.transform = 'none'; el.style.width = '{width}px'; "
                f"el.style.height = '{height}px'; window.dispatchEvent(new Event('resize')); return true; }})()")
        for sel in args.get("hideSelectors") or []:
            await pv.eval(f"document.querySelectorAll({json.dumps(str(sel))}).forEach(e => e.style.display = 'none')")
        await pv._page.set_viewport_size({"width": width, "height": height})
        notes: list[str] = []
        try:
            raw_notes = await pv.eval("(() => { const s = document.getElementById('speaker-notes'); "
                                      "return s ? JSON.parse(s.textContent) : null; })()")
            if isinstance(raw_notes, list):
                notes = [str(n) for n in raw_notes]
        except Exception:
            pass
        captured: list[dict[str, Any]] = []
        for s in slides:
            if s.get("showJs"):
                await pv.eval(str(s["showJs"]))
            await pv._page.wait_for_timeout(max(0, min(int(s.get("delay", 600)), 10_000)))
            entry: dict[str, Any] = {}
            dom = await pv._page.evaluate(_SLIDE_DOM_JS, str(s["selector"]))
            if dom is None:
                return err(f"selector {s['selector']!r} matched nothing on the page.")
            entry["box"] = dom["box"]
            entry["items"] = dom["items"] if mode == "editable" else []
            entry["background"] = bool(args.get("background"))
            if mode == "screenshots" or args.get("background"):
                el = await pv._page.query_selector(str(s["selector"]))
                entry["png"] = await (el.screenshot(type="png") if el else pv._page.screenshot(type="png"))
            captured.append(entry)
        data, flags = await in_thread(_build_pptx, width, height, captured, notes, mode,
                                      ctx().workspace.resolve())
    except Exception as e:
        return err(f"gen_pptx failed: {type(e).__name__}: {e}")
    finally:
        try:
            from ..gui.preview import VIEWPORT

            await pv._page.set_viewport_size(dict(VIEWPORT))
        except Exception:
            pass
    out.write_bytes(data)
    return ok(f"Wrote {out} ({len(data):,} bytes, {len(slides)} slide(s), {mode} mode)."
              + ("\nValidation flags:\n" + "\n".join(f"- {f}" for f in flags) if flags else "\nNo validation flags."))


# --- open_for_print ---------------------------------------------------------------------------


@tool(
    "open_for_print",
    "Print an HTML file to PDF from the hidden frame and return the PDF's path. For a "
    "deck built on the deck starter, pass width=1920 and height=1080 to get one page per "
    "slide at canvas size; for a document, omit them and the page's own @page rule (or "
    "Letter) applies. Backgrounds are printed.",
    {"type": "object", "properties": {
        "project_relative_file_path": {"type": "string", "description": "The HTML file to print."},
        "output_path": {"type": "string", "description": "Where to write the PDF (default: beside the HTML)."},
        "width": {"type": "number"}, "height": {"type": "number"}, "landscape": {"type": "boolean"}},
     "required": ["project_relative_file_path"]},
)
async def open_for_print(args: dict[str, Any]) -> dict[str, Any]:
    raw = args.get("project_relative_file_path") or args.get("path")
    if not raw:
        return err("open_for_print needs 'project_relative_file_path'.")
    p = _resolve(str(raw))
    if not p.is_file():
        return err(f"No file at {p}")
    out = _resolve(str(args.get("output_path") or "")) if args.get("output_path") else p.with_suffix(".pdf")
    if not _inside(out):
        return err("output_path must be inside the workspace.")
    pv = get_preview()
    try:
        if pv.loaded != p.resolve():
            await pv.load(p)
        w, h = args.get("width"), args.get("height")
        out.parent.mkdir(parents=True, exist_ok=True)
        await pv.pdf(out, width=int(w) if w else None, height=int(h) if h else None,
                     landscape=bool(args.get("landscape")))
    except Exception as e:
        return err(f"Could not print: {type(e).__name__}: {e}")
    return ok(f"Wrote {out} ({out.stat().st_size:,} bytes). Open it, or hand it to the user.")


EXPORT_TOOLS = [super_inline_html, gen_pptx, open_for_print]
