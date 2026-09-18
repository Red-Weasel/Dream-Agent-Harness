"""Structural handoff checks must survive clean preview logs and absent Studio."""
import io
import os
import zipfile
from unittest.mock import AsyncMock

import pytest

from dream.tools import context, studio
from dream.tools.context import ToolContext
from dream.workflows.validation import verify


@pytest.fixture
def handoff(tmp_path, monkeypatch):
    monkeypatch.setattr(context, '_CTX', ToolContext(
        store=None, working=None, browser=None, session_id='artifact-check',
        workspace=tmp_path))
    monkeypatch.setattr(studio, 'studio', lambda: None)
    load = AsyncMock(return_value=[])
    monkeypatch.setattr(studio, '_load', load)
    return tmp_path, load


@pytest.mark.parametrize('tag', ['img', 'audio', 'video', 'source', 'script', 'track'])
async def test_missing_local_dependency_cannot_earn_clean_handoff(handoff, tag):
    root, load = handoff
    (root / 'page.html').write_text(f'<h1>Finished</h1><{tag} src="missing.bin">')
    result = await studio.done.handler({'path': 'page.html'})
    assert result.get('is_error'), result
    assert 'missing.bin' in result['content'][0]['text']
    assert 'call done again' in result['content'][0]['text']
    load.assert_not_awaited()


@pytest.mark.parametrize('reference', ['../secret.png', '%2e%2e/secret.png', '/etc/passwd', 'file:///etc/passwd', 'alias.png'])
async def test_dependency_path_escape_is_rejected_before_preview(handoff, reference):
    root, load = handoff
    (root / 'alias.png').symlink_to('/etc/passwd')
    (root / 'page.html').write_text(f'<img src="{reference}">')
    result = await studio.done.handler({'path': 'page.html'})
    assert result.get('is_error'), result
    load.assert_not_awaited()


@pytest.mark.parametrize('damage', ['empty', 'fifo'])
async def test_unusable_local_dependency_is_rejected_without_blocking(handoff, damage):
    root, load = handoff
    asset = root / 'asset.bin'
    if damage == 'fifo':
        os.mkfifo(asset)
    else:
        asset.touch()
    (root / 'page.html').write_text('<img src="asset.bin">')
    assert (await studio.done.handler({'path': 'page.html'})).get('is_error')
    load.assert_not_awaited()


async def test_large_regular_media_keeps_presence_and_decode_claims_separate(handoff):
    root, load = handoff
    with (root / 'clip.mp4').open('wb') as stream:
        stream.truncate(64 * 1024 * 1024)
    (root / 'page.html').write_text('<video src="clip.mp4"></video>')
    result = await studio.done.handler({'path': 'page.html'})
    assert not result.get('is_error'), result
    assert 'media decoding' in result['content'][0]['text']
    assert 'unverified' in result['content'][0]['text']
    load.assert_awaited_once()


async def test_existing_relative_dependency_accepts_without_claiming_quality(handoff):
    root, load = handoff
    (root / 'asset.js').write_text('window.answer = 42;')
    (root / 'page.html').write_text('<h1>Ready</h1><script src="asset.js?v=1#code"></script>')
    result = await studio.done.handler({'path': 'page.html'})
    assert not result.get('is_error'), result
    message = result['content'][0]['text'].lower()
    assert 'local' in message and 'unverified' in message and 'visual' in message
    load.assert_awaited_once()


async def test_parent_relative_dependency_inside_workspace_is_preserved(handoff):
    root, load = handoff
    (root / 'pages').mkdir()
    (root / 'assets').mkdir()
    (root / 'assets/image.svg').write_text('<svg/>')
    (root / 'pages/page.html').write_text('<img src="../assets/image.svg">')
    assert not (await studio.done.handler({'path': 'pages/page.html'})).get('is_error')
    load.assert_awaited_once()


@pytest.mark.parametrize('markup,failed', [
    ('<img src="missing.png" src="asset.svg">', True),
    ('<img src="asset.svg" src="missing.png">', False),
    ('<link rel="StyleSheet" href="missing.css">', True),
])
async def test_static_references_match_browser_attribute_semantics(handoff, markup, failed):
    root, load = handoff
    (root / 'asset.svg').write_text('<svg/>')
    (root / 'page.html').write_text(markup)
    result = await studio.done.handler({'path': 'page.html'})
    assert bool(result.get('is_error')) is failed
    assert load.await_count == (0 if failed else 1)


@pytest.mark.parametrize('body', ['', '   ', '\x00'])
async def test_empty_or_binary_page_is_not_a_deliverable(handoff, body):
    root, load = handoff
    (root / 'page.html').write_text(body)
    assert (await studio.done.handler({'path': 'page.html'})).get('is_error')
    load.assert_not_awaited()


async def test_remote_references_remain_explicitly_unverified(handoff):
    root, _ = handoff
    (root / 'page.html').write_text('<img src="https://example.invalid/picture.png">')
    result = await studio.done.handler({'path': 'page.html'})
    assert not result.get('is_error'), result
    message = result['content'][0]['text'].lower()
    assert 'remote' in message and 'unverified' in message


@pytest.mark.parametrize('tag', ['<video poster="missing.png"></video>',
                                '<link rel="stylesheet" href="missing.css">',
                                '<svg><image href="missing.svg"/></svg>'])
async def test_other_static_dependency_attributes_are_checked(handoff, tag):
    root, load = handoff
    (root / 'page.html').write_text(tag)
    assert (await studio.done.handler({'path': 'page.html'})).get('is_error')
    load.assert_not_awaited()


async def test_too_many_dependencies_fail_before_preview(handoff):
    root, load = handoff
    (root / 'page.html').write_text('<img src="data:image/png;base64,AA==">' * 129)
    result = await studio.done.handler({'path': 'page.html'})
    assert result.get('is_error') and '128' in result['content'][0]['text']
    load.assert_not_awaited()


async def test_page_symlink_rejected_before_preview(handoff):
    root, load = handoff
    (root / 'real.html').write_text('<h1>Ready</h1>')
    (root / 'page.html').symlink_to(root / 'real.html')
    assert (await studio.done.handler({'path': 'page.html'})).get('is_error')
    load.assert_not_awaited()


async def test_inline_svg_and_script_strings_are_not_false_dependencies(handoff):
    root, _ = handoff
    (root / 'page.html').write_text('<svg><use href="#icon"/></svg>'
                                  '<script>const sample = `<img src="missing.png">`;</script>')
    assert not (await studio.done.handler({'path': 'page.html'})).get('is_error')


def _deck():
    from pptx import Presentation
    from PIL import Image
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[0])
    slide.shapes.title.text = 'Observed findings'
    bitmap = io.BytesIO()
    Image.new('RGB', (4, 4), 'green').save(bitmap, format='PNG')
    bitmap.seek(0)
    slide.shapes.add_picture(bitmap, 0, 0)
    output = io.BytesIO()
    deck.save(output)
    return output.getvalue()


def _rewrite(data, *, omit=None, extra=None):
    output = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(data)) as source, zipfile.ZipFile(output, 'w') as target:
        for entry in source.infolist():
            if entry.filename != omit:
                target.writestr(entry, source.read(entry))
        if extra:
            target.writestr(*extra)
    return output.getvalue()


def test_presentation_missing_linked_media_is_not_verified():
    with pytest.raises(ValueError, match='missing'):
        verify(_rewrite(_deck(), omit='ppt/media/image1.png'), 'presentation')


async def test_missing_presentation_media_blocks_runtime_review_state(tmp_path):
    from dream.workflows import WorkflowService
    service = WorkflowService(tmp_path)
    task = service.create('presentation', {'goal': 'Summarize fixture observations'})
    await service.start(task['id'], task['version'], 'fixture-start', lambda *_: None)
    output = tmp_path / task['output_path']
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_bytes(_rewrite(_deck(), omit='ppt/media/image1.png'))
    result = service.event(task['id'], 'final', attempt=1)
    assert result['status'] == 'needs_attention'
    assert 'missing' in result['message'] and result['artifacts'] == []


async def test_valid_nested_absolute_page_and_asset_path_are_accepted(handoff):
    root, _ = handoff
    folder = root / 'pages'
    folder.mkdir()
    (folder / 'image.svg').write_text('<svg xmlns="http://www.w3.org/2000/svg"/>')
    page = folder / 'page.html'
    page.write_text('<img src="image.svg">')
    assert not (await studio.done.handler({'path': str(page)})).get('is_error')


@pytest.mark.parametrize('name', ['../outside.xml', '/outside.xml', 'ppt/../outside.xml'])
def test_unsafe_presentation_members_are_rejected(name):
    with pytest.raises(ValueError, match='unsafe'):
        verify(_rewrite(_deck(), extra=(name, b'<unused/>')), 'presentation')


def test_duplicate_presentation_members_are_rejected():
    original = _deck()
    with zipfile.ZipFile(io.BytesIO(original)) as package:
        slide = package.read('ppt/slides/slide1.xml')
    with pytest.warns(UserWarning, match='Duplicate'):
        data = _rewrite(original, extra=('ppt/slides/slide1.xml', slide))
    with pytest.raises(ValueError, match='duplicate'):
        verify(data, 'presentation')


def test_valid_presentation_keeps_factual_and_visual_quality_for_review():
    result = verify(_deck(), 'presentation')
    assert 'slide' in result and 'visual quality still need your review' in result
