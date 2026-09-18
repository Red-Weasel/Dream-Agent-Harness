"""super_inline_html and gen_pptx, against a deck built from the starter."""

from __future__ import annotations

import io
import shutil
from pathlib import Path

import pytest
from pptx import Presentation

from dream import config
from dream.gui import preview as preview_mod
from dream.tools import context as tool_context
from dream.tools.context import ToolContext, set_context
from dream.tools.export_tools import EXPORT_TOOLS, gen_pptx, super_inline_html

STARTER = Path(__file__).parent.parent / "dream" / "gui" / "starters" / "deck_stage.js"

DECK = """<!doctype html><html><head><meta charset="utf-8">
<template id="__bundler_thumbnail"><svg viewBox="0 0 10 10"><rect width="10" height="10" fill="#0af"/></svg></template>
<script type="application/json" id="speaker-notes">["Say hello", "Show the plan", ""]</script>
<link rel="stylesheet" href="deck.css">
</head><body>
<deck-stage>
  <section><h1 id="t1">One</h1><p>first slide</p></section>
  <section style="background:#123"><h1 style="color:#fff">Two</h1><img src="dot.png" width="40" height="40"></section>
  <section><h2>Three</h2></section>
</deck-stage>
<script src="deck_stage.js"></script>
</body></html>"""


@pytest.fixture
async def ws(tmp_path, monkeypatch):
    set_context(ToolContext(store=None, working=None, browser=None,  # type: ignore[arg-type]
                            session_id="t", workspace=tmp_path))
    monkeypatch.setattr(config, "SCREENSHOT_DIR", tmp_path / "var" / "shots")
    preview_mod._PREVIEW = None
    shutil.copy(STARTER, tmp_path / "deck_stage.js")
    (tmp_path / "deck.css").write_text("section{font:64px system-ui;padding:80px;background:#fff}")
    from PIL import Image
    Image.new("RGB", (40, 40), (255, 0, 0)).save(tmp_path / "dot.png")
    (tmp_path / "deck.html").write_text(DECK)
    yield tmp_path
    if preview_mod._PREVIEW is not None:
        await preview_mod._PREVIEW.aclose()
        preview_mod._PREVIEW = None
    tool_context._CTX = None


def _text(res):
    return res["content"][0]["text"]


def _failed(res):
    return bool(res.get("is_error"))


@pytest.mark.asyncio
async def test_super_inline_html_writes_one_offline_file(ws):
    res = await super_inline_html.handler({"input_path": "deck.html", "output_path": "out/deck-standalone.html"})
    assert not _failed(res), _text(res)
    out = (ws / "out" / "deck-standalone.html").read_text()
    assert "customElements.define('deck-stage'" in out
    assert '<script src="deck_stage.js"></script>' not in out, "the tag was replaced by the inlined code"
    assert "data:image/png;base64," in out and "<style>section{font:64px" in out
    bad = await super_inline_html.handler({"input_path": "deck.html", "output_path": "../x.html"})
    assert _failed(bad) and "inside the workspace" in _text(bad)
    (ws / "plain.html").write_text("<p>no splash</p>")
    bad = await super_inline_html.handler({"input_path": "plain.html", "output_path": "p.html"})
    assert _failed(bad) and "__bundler_thumbnail" in _text(bad)


@pytest.mark.asyncio
async def test_gen_pptx_screenshots_mode_one_picture_per_slide_with_notes(ws):
    res = await gen_pptx.handler({
        "path": "deck.html", "width": 1920, "height": 1080, "mode": "screenshots",
        "resetTransformSelector": "deck-stage", "hideSelectors": [],
        "slides": [{"selector": "deck-stage > section:nth-of-type(%d)" % (i + 1), "showJs": f"goToSlide({i})", "delay": 80}
                   for i in range(3)],
    })
    assert not _failed(res), _text(res)
    prs = Presentation(str(ws / "deck.pptx"))
    assert len(prs.slides) == 3
    assert (prs.slide_width, prs.slide_height) == (1920 * 9525, 1080 * 9525)
    for slide in prs.slides:
        assert any(sh.shape_type == 13 for sh in slide.shapes), "a picture per slide"
    assert prs.slides[0].notes_slide.notes_text_frame.text == "Say hello"
    assert prs.slides[1].notes_slide.notes_text_frame.text == "Show the plan"
    assert "duplicate_adjacent" not in _text(res) and "slide_size_mismatch" not in _text(res)


@pytest.mark.asyncio
async def test_gen_pptx_editable_mode_emits_text_and_pictures_and_flags_bad_captures(ws):
    res = await gen_pptx.handler({
        "path": "deck.html", "width": 1920, "height": 1080, "mode": "editable", "filename": "edit deck",
        "resetTransformSelector": "deck-stage",
        "slides": [{"selector": "deck-stage > section:nth-of-type(1)", "showJs": "goToSlide(0)", "delay": 60},
                   {"selector": "deck-stage > section:nth-of-type(2)", "showJs": "goToSlide(1)", "delay": 60}],
    })
    assert not _failed(res), _text(res)
    prs = Presentation(str(ws / "edit_deck.pptx"))
    texts = [sh.text_frame.text for sh in prs.slides[0].shapes if sh.has_text_frame]
    assert "One" in texts and "first slide" in texts
    assert any(sh.shape_type == 13 for sh in prs.slides[1].shapes), "the image became a picture"
    assert any(sh.shape_type == 1 for sh in prs.slides[1].shapes), "the colored section became a shape"
    # no showJs → the same slide twice → flagged; a tiny selector → size mismatch
    res = await gen_pptx.handler({
        "path": "deck.html", "width": 1920, "height": 1080, "mode": "screenshots", "filename": "bad",
        "resetTransformSelector": "deck-stage",
        "slides": [{"selector": "deck-stage > section[data-active]", "showJs": "goToSlide(0)", "delay": 60},
                   {"selector": "deck-stage > section[data-active]", "delay": 60},
                   {"selector": "#t1", "delay": 20}],
    })
    assert not _failed(res)
    assert "duplicate_adjacent: slide 2" in _text(res) and "slide_size_mismatch: slide 3" in _text(res)
    bad = await gen_pptx.handler({"path": "deck.html", "width": 1920, "height": 1080,
                                  "slides": [{"selector": "#nope"}]})
    assert _failed(bad) and "matched nothing" in _text(bad)


def test_the_three_tools_are_exported():
    assert [t.name for t in EXPORT_TOOLS] == ["super_inline_html", "gen_pptx", "open_for_print"]


def test_super_inline_html_says_what_it_leaves_alone():
    """The bundle is bounded by the workspace, not the page's folder — the
    description must say so."""
    assert "outside the workspace" in super_inline_html.description
    assert "page's folder" not in super_inline_html.description


@pytest.mark.asyncio
async def test_editable_mode_never_embeds_a_file_image_from_outside_the_workspace(ws):
    """Gate 7 observation: the DOM walker drew every <img> to a canvas, so a
    file:// image outside the workspace arrived as a data URI and bypassed the
    workspace bound that _image_bytes enforces."""
    from PIL import Image
    Image.new("RGB", (30, 30), (0, 255, 0)).save(ws.parent / "outside.png")
    (ws / "leak.html").write_text(DECK.replace('src="dot.png"', 'src="../outside.png"'))
    res = await gen_pptx.handler({
        "path": "leak.html", "width": 1920, "height": 1080, "mode": "editable", "filename": "leak",
        "resetTransformSelector": "deck-stage",
        "slides": [{"selector": "deck-stage > section:nth-of-type(2)", "showJs": "goToSlide(1)", "delay": 60}],
    })
    assert not _failed(res), _text(res)
    assert "image_skipped: slide 1" in _text(res)
    prs = Presentation(str(ws / "leak.pptx"))
    assert not any(sh.shape_type == 13 for sh in prs.slides[0].shapes), "no picture from outside"
