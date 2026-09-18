"""Real file-reading fixtures and hostile document regressions, without models."""
from __future__ import annotations

import io
import json
from pathlib import Path
import shutil
import struct
import tarfile
import zipfile

import pytest
from PIL import Image

from dream.tools.context import ToolContext, bind_context
from dream.tools.native import read_file


@pytest.fixture
def workspace(tmp_path):
    with bind_context(ToolContext(store=None, working=None, browser=None,
                                  session_id='documents', workspace=tmp_path)):
        yield tmp_path


def text(result):
    return result['content'][0]['text']


def office(path, entries):
    with zipfile.ZipFile(path, 'w', compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr('[Content_Types].xml', '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"/>')
        for name, data in entries.items():
            archive.writestr(name, data)


def pdf(path, pages):
    """A valid PDF with real font/content streams, offsets and page tree."""
    objects = [b'<< /Type /Catalog /Pages 2 0 R >>', b'',
               b'<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>']
    page_ids = []
    for value in pages:
        page_id = len(objects) + 1
        page_ids.append(page_id)
        objects.append(f'<< /Type /Page /Parent 2 0 R /MediaBox [0 0 400 300] /Resources << /Font << /F1 3 0 R >> >> /Contents {page_id+1} 0 R >>'.encode())
        stream = f'BT /F1 14 Tf 30 250 Td ({value}) Tj ET'.encode()
        objects.append(b'<< /Length ' + str(len(stream)).encode() + b' >>\nstream\n' + stream + b'\nendstream')
    objects[1] = f'<< /Type /Pages /Count {len(page_ids)} /Kids [{" ".join(str(i)+" 0 R" for i in page_ids)}] >>'.encode()
    data = bytearray(b'%PDF-1.4\n')
    offsets = [0]
    for i, value in enumerate(objects, 1):
        offsets.append(len(data)); data.extend(f'{i} 0 obj\n'.encode() + value + b'\nendobj\n')
    xref = len(data)
    data.extend(f'xref\n0 {len(objects)+1}\n0000000000 65535 f \n'.encode())
    for offset in offsets[1:]:
        data.extend(f'{offset:010d} 00000 n \n'.encode())
    data.extend(f'trailer << /Size {len(objects)+1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n'.encode())
    path.write_bytes(data)


@pytest.mark.asyncio
async def test_docx_extracts_paragraphs_and_tables_without_zip_garbage(workspace):
    office(workspace/'note.docx', {'word/document.xml': '''<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body>
    <w:p><w:r><w:t>Quarterly </w:t></w:r><w:r><w:t>report</w:t></w:r></w:p>
    <w:tbl><w:tr><w:tc><w:p><w:r><w:t>Revenue</w:t></w:r></w:p></w:tc><w:tc><w:p><w:r><w:t>42</w:t></w:r></w:p></w:tc></w:tr></w:tbl>
    </w:body></w:document>'''})
    result = await read_file.handler({'path': 'note.docx'})
    assert not result.get('is_error'), text(result)
    assert 'Quarterly report' in text(result)
    assert 'Revenue | 42' in text(result)
    assert 'layout' in text(result).lower()


@pytest.mark.asyncio
async def test_binary_rejected_and_pagination_is_actionable(workspace):
    (workspace/'program.bin').write_bytes(b'\x7fELF\x00\x01binary')
    result = await read_file.handler({'path': 'program.bin'})
    assert result.get('is_error')
    assert 'binary' in text(result).lower()
    (workspace/'lines.txt').write_text('abcdefghij')
    result = await read_file.handler({'path': 'lines.txt', 'offset': 3, 'limit': 4})
    assert 'defg' in text(result) and 'offset=7' in text(result)
    assert 'abcdefghij' not in text(result)


@pytest.mark.asyncio
async def test_pdf_page_selection_and_blank_page_report(workspace):
    if not shutil.which('pdfinfo') or not shutil.which('pdftotext'):
        pytest.skip('Poppler unavailable')
    pdf(workspace/'report.pdf', ['First page text', 'Second page text', ''])
    result = await read_file.handler({'path': 'report.pdf', 'page_start': 2, 'page_count': 2})
    assert not result.get('is_error'), text(result)
    assert '[Page 2]' in text(result) and 'Second page text' in text(result)
    assert 'First page text' not in text(result)
    assert 'scanned or blank' in text(result)


def workbook(path, target='worksheets/sheet2.xml', external=False):
    mode = ' TargetMode="External"' if external else ''
    office(path, {
        'xl/workbook.xml': '<workbook xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"><sheets><sheet name="Revenue" sheetId="7" r:id="rId9"/><sheet name="Hidden notes" sheetId="8" r:id="rId10" state="hidden"/></sheets></workbook>',
        'xl/_rels/workbook.xml.rels': f'<Relationships><Relationship Id="rId9" Target="{target}"{mode}/><Relationship Id="rId10" Target="worksheets/sheet1.xml"/></Relationships>',
        'xl/sharedStrings.xml': '<sst><si><r><t>North </t></r><r><t>team</t></r></si></sst>',
        'xl/worksheets/sheet2.xml': '<worksheet><sheetData><row r="1"><c r="A1" t="s"><v>0</v></c><c r="B1"><v>42</v></c><c r="C1"><f>B1*2</f><v>84</v></c><c r="D1"><f>NOW()</f></c><c r="E1" t="b"><v>1</v></c></row></sheetData></worksheet>',
        'xl/worksheets/sheet1.xml': '<worksheet><sheetData><row><c r="A1" t="inlineStr"><is><t>Private sheet note</t></is></c></row></sheetData></worksheet>',
    })


@pytest.mark.asyncio
async def test_xlsx_sheet_relationships_strings_formulas_and_hidden_sheets(workspace):
    workbook(workspace/'book.xlsx')
    result = await read_file.handler({'path': 'book.xlsx', 'sheet': 'Revenue'})
    assert not result.get('is_error'), text(result)
    output = text(result)
    assert 'A1\t"North team"' in output
    assert 'B1\t"42"' in output
    assert '=B1*2 [cached: 84]' in output and '=NOW() [cached value unavailable]' in output
    assert 'E1\t"TRUE"' in output
    assert 'Private sheet note' not in output
    result = await read_file.handler({'path': 'book.xlsx'})
    assert 'Private sheet note' in text(result) and 'hidden' in text(result)
    result = await read_file.handler({'path': 'book.xlsx', 'sheet': 'Typo'})
    assert result.get('is_error') and 'Available sheets' in text(result)


@pytest.mark.asyncio
async def test_real_pptx_text_order_and_notes(workspace):
    from pptx import Presentation
    deck = Presentation()
    for title in ('First topic', 'Second topic'):
        slide = deck.slides.add_slide(deck.slide_layouts[1])
        slide.shapes.title.text = title
        slide.placeholders[1].text = 'Bullet text'
        slide.notes_slide.notes_text_frame.text = title + ' speaker explanation'
    # Reorder the presentation independently of numbered filenames.
    deck.slides._sldIdLst.insert(0, deck.slides._sldIdLst[1])
    deck.save(workspace/'deck.pptx')
    result = await read_file.handler({'path': 'deck.pptx'})
    assert not result.get('is_error'), text(result)
    output = text(result)
    assert output.index('Second topic') < output.index('First topic')
    assert '[Slide 1]' in output and '[Slide 2]' in output and '[Speaker notes]' in output
    assert 'Second topic speaker explanation' in output


@pytest.mark.asyncio
@pytest.mark.parametrize('name', ['code.py', 'table.csv', 'data.json', 'page.html', 'data.xml', 'vector.svg'])
async def test_text_formats_preserve_source(workspace, name):
    original = '<tag>héllo</tag>\nvalue,42\n'
    (workspace/name).write_text(original)
    result = await read_file.handler({'path': name})
    assert not result.get('is_error') and text(result) == original


@pytest.mark.asyncio
@pytest.mark.parametrize('encoding', ['utf-8-sig', 'utf-16', 'utf-32'])
async def test_bom_unicode_is_decoded_without_replacement(workspace, encoding):
    (workspace/'unicode.txt').write_bytes('Résumé 日本語\n'.encode(encoding))
    result = await read_file.handler({'path': 'unicode.txt'})
    assert text(result) == 'Résumé 日本語\n'


@pytest.mark.asyncio
async def test_images_are_metadata_not_visual_understanding(workspace):
    Image.new('RGB', (40, 30), 'blue').save(workspace/'photo.png')
    result = await read_file.handler({'path': 'photo.png'})
    assert not result.get('is_error'), text(result)
    assert 'size: 40×30' in text(result) and 'metadata only' in text(result)
    assert 'blue' not in text(result)
    (workspace/'corrupt.png').write_bytes(b'not a picture')
    assert (await read_file.handler({'path': 'corrupt.png'})).get('is_error')


@pytest.mark.asyncio
async def test_zip_and_compressed_tar_are_listed_without_extracting(workspace):
    with zipfile.ZipFile(workspace/'bundle.zip', 'w') as archive:
        archive.writestr('folder/note.txt', 'secret member contents')
    result = await read_file.handler({'path': 'bundle.zip'})
    assert not result.get('is_error'), text(result)
    assert 'folder/note.txt' in text(result) and 'secret member contents' not in text(result)
    with tarfile.open(workspace/'bundle.tar.gz', 'w:gz') as archive:
        member = tarfile.TarInfo('note.txt'); member.size = 6
        archive.addfile(member, io.BytesIO(b'secret'))
        link = tarfile.TarInfo('link'); link.type = tarfile.SYMTYPE; link.linkname = '/etc/passwd'
        archive.addfile(link)
    result = await read_file.handler({'path': 'bundle.tar.gz'})
    assert not result.get('is_error'), text(result)
    assert 'link (not followed)' in text(result) and 'secret' not in text(result)
    assert not (workspace/'folder').exists() and not (workspace/'note.txt').exists()


@pytest.mark.asyncio
@pytest.mark.parametrize('member', ['../outside.txt', '/tmp/outside.txt', 'C:/outside.txt', 'folder\\outside.txt'])
async def test_archive_traversal_is_refused(workspace, member):
    with zipfile.ZipFile(workspace/'hostile.zip', 'w') as archive:
        archive.writestr(member, 'never extract')
    result = await read_file.handler({'path': 'hostile.zip'})
    assert result.get('is_error') and 'Unsafe archive member path' in text(result)


@pytest.mark.asyncio
async def test_office_bomb_and_duplicate_parts_are_refused(workspace):
    office(workspace/'bomb.docx', {'word/document.xml': '<document>' + 'a'*500_000 + '</document>'})
    result = await read_file.handler({'path': 'bomb.docx'})
    assert result.get('is_error') and 'compression-ratio' in text(result)
    office(workspace/'duplicate.docx', {'word/document.xml': '<document/>'})
    with pytest.warns(UserWarning, match='Duplicate'):
        with zipfile.ZipFile(workspace/'duplicate.docx', 'a') as archive:
            archive.writestr('word/document.xml', '<document/>')
    result = await read_file.handler({'path': 'duplicate.docx'})
    assert result.get('is_error') and 'duplicate ZIP' in text(result)


@pytest.mark.asyncio
@pytest.mark.parametrize('encoding', ['utf-8', 'utf-16'])
async def test_office_dtd_entities_are_never_expanded(workspace, encoding):
    office(workspace/'entity.docx', {'word/document.xml': '<!DOCTYPE document [<!ENTITY x SYSTEM "file:///etc/passwd">]><document>&x;</document>'.encode(encoding)})
    result = await read_file.handler({'path': 'entity.docx'})
    assert result.get('is_error') and 'DTD or entity' in text(result)
    assert 'root:x' not in text(result)


@pytest.mark.asyncio
@pytest.mark.parametrize('target,external', [('http://example.test/secret', True), ('../../../etc/passwd', False)])
async def test_workbook_external_or_escaping_relationship_is_refused(workspace, target, external):
    workbook(workspace/'external.xlsx', target, external)
    result = await read_file.handler({'path': 'external.xlsx', 'sheet': 'Revenue'})
    assert result.get('is_error')
    assert 'external relationship' in text(result) or 'Unsafe archive member' in text(result)


@pytest.mark.asyncio
async def test_scanned_pdf_corrupt_pdf_and_missing_dependency(workspace, monkeypatch):
    Image.new('RGB', (80, 60), 'white').save(workspace/'scan.pdf', format='PDF')
    result = await read_file.handler({'path': 'scan.pdf'})
    assert result.get('is_error') and 'OCR has not been run' in text(result)
    (workspace/'corrupt.pdf').write_bytes(b'%PDF-1.7\nnot a PDF')
    result = await read_file.handler({'path': 'corrupt.pdf'})
    assert result.get('is_error') and 'corrupt or restricted' in text(result)
    monkeypatch.setattr('dream.files.reading.shutil.which', lambda _: None)
    result = await read_file.handler({'path': 'scan.pdf'})
    assert result.get('is_error') and 'requires Poppler' in text(result)


@pytest.mark.asyncio
async def test_encrypted_office_and_unsupported_media_are_actionable(workspace):
    (workspace/'protected.docx').write_bytes(b'\xd0\xcf\x11\xe0' + b'\0'*100)
    result = await read_file.handler({'path': 'protected.docx'})
    assert result.get('is_error') and 'unencrypted' in text(result)
    for suffix in ('doc', 'xls', 'ppt', 'docm', 'mp3', 'mp4', '7z'):
        name = 'unsupported.' + suffix; (workspace/name).write_bytes(b'\0binary')
        result = await read_file.handler({'path': name})
        assert result.get('is_error')
        assert 'transcribed' in text(result) if suffix in ('mp3', 'mp4') else 'supported' in text(result)


@pytest.mark.asyncio
async def test_limits_report_partial_content_and_reject_invalid_options(workspace, monkeypatch):
    import dream.files.reading as reading
    (workspace/'large.txt').write_text('x'*(reading.MAX_EXTRACT_CHARS + 1))
    result = await read_file.handler({'path': 'large.txt', 'limit': 100})
    assert 'offset=100' in text(result) and 'not the complete file' in text(result)
    for options in ({'limit': 50001}, {'limit': True}, {'offset': -1}, {'page_count': 51}, {'sheet': ''}, {'page_start': 2}):
        result = await read_file.handler({'path': 'large.txt', **options})
        assert result.get('is_error'), options
    monkeypatch.setattr(reading, 'MAX_FILE_BYTES', 10)
    result = await read_file.handler({'path': 'large.txt'})
    assert result.get('is_error') and 'reading limit' in text(result)


@pytest.mark.asyncio
async def test_workbook_cell_and_tar_decompression_limits(workspace, monkeypatch):
    import dream.files.reading as reading
    workbook(workspace/'cells.xlsx')
    monkeypatch.setattr(reading, 'MAX_CELLS', 2)
    result = await read_file.handler({'path': 'cells.xlsx'})
    assert not result.get('is_error')
    assert 'A1' in text(result) and 'B1' in text(result) and 'C1' not in text(result)
    assert 'not the complete file' in text(result)
    with tarfile.open(workspace/'large.tar.gz', 'w:gz') as archive:
        item = tarfile.TarInfo('big'); item.size = 10000
        archive.addfile(item, io.BytesIO(b'0'*10000))
    monkeypatch.setattr(reading, 'MAX_EXPANDED_BYTES', 1024)
    result = await read_file.handler({'path': 'large.tar.gz'})
    assert result.get('is_error') and 'expanded-size' in text(result)


@pytest.mark.asyncio
async def test_workspace_relative_and_existing_absolute_read_semantics(workspace, tmp_path_factory):
    from dream.core import policy
    external = tmp_path_factory.mktemp('document-external')/'permitted.txt'
    external.write_text('Allowed read-only file')
    (workspace/'reference.txt').symlink_to(external)
    for path in (str(external), 'reference.txt'):
        assert policy.decide('read_file', {'path': path}, 'auto', workspace)[0] == 'allow'
        result = await read_file.handler({'path': path})
        assert text(result) == 'Allowed read-only file'


def test_pdf_process_limits_and_password_error(monkeypatch):
    import sys
    import time
    import dream.files.reading as reading
    monkeypatch.setattr(reading, 'MAX_TEXT_BYTES', 1000)
    with pytest.raises(reading.DocumentError, match='output exceeded'):
        reading._process([sys.executable, '-c', 'print("x"*10000)'])
    monkeypatch.setattr(reading, 'PDF_TIMEOUT', .1)
    started = time.monotonic()
    with pytest.raises(reading.DocumentError, match='timed out'):
        reading._process([sys.executable, '-c', 'import time; time.sleep(10)'])
    assert time.monotonic() - started < 2
    with pytest.raises(reading.DocumentError, match='password-protected'):
        reading._process([sys.executable, '-c', 'import sys; sys.stderr.write("Incorrect password"); sys.exit(1)'])


@pytest.mark.parametrize('declared', [1, 5000, 65535])
@pytest.mark.parametrize('suffix', ['zip', 'docx'])
def test_zip_member_precheck_runs_before_zipfile_allocation(workspace, monkeypatch, declared, suffix):
    import dream.files.reading as reading
    path = workspace/('many.' + suffix)
    with zipfile.ZipFile(path, 'w') as archive:
        for index in range(3):
            archive.writestr(str(index), b'')
    data = bytearray(path.read_bytes())
    eocd = data.rfind(b'PK\x05\x06')
    struct.pack_into('<HH', data, eocd + 8, declared, declared)
    path.write_bytes(data)
    monkeypatch.setattr(reading, 'MAX_ZIP_MEMBERS', 2)
    def forbidden(*args, **kwargs):
        raise AssertionError('ZipFile constructor allocated before the entry-count precheck')
    monkeypatch.setattr(reading.zipfile, 'ZipFile', forbidden)
    with pytest.raises(reading.DocumentError, match='members|ZIP64'):
        reading.read_document(path)
