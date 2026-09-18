"""CPU-only durable workflow and artifact acceptance tests."""
import asyncio
import importlib.util
import json
from pathlib import Path
import pytest


def service(tmp_path):
    assert importlib.util.find_spec('dream.workflows') is not None, 'Guided workflow service is missing'
    from dream.workflows import WorkflowService
    return WorkflowService(tmp_path)


def draft(svc, **extra):
    return svc.create('report', {'goal': 'Explain the observations', **extra})


def test_review_then_idempotent_start_and_version_conflict(tmp_path):
    svc = service(tmp_path)
    task = draft(svc)
    assert task['status'] == 'draft'
    calls = []
    async def dispatch(prompt, metadata):
        calls.append(metadata)
    first = asyncio.run(svc.start(task['id'], task['version'], 'request-one', dispatch))
    again = asyncio.run(svc.start(task['id'], task['version'], 'request-one', dispatch))
    assert first['status'] == again['status'] == 'queued'
    assert len(calls) == 1 and calls[0]['workflow_task_id'] == task['id']
    with pytest.raises(ValueError, match='changed'):
        asyncio.run(svc.start(task['id'], task['version'], 'request-two', dispatch))


def test_missing_output_cannot_be_claimed_complete(tmp_path):
    svc = service(tmp_path); task = draft(svc)
    task = asyncio.run(svc.start(task['id'], task['version'], 'one', lambda *_: None))
    result = svc.event(task['id'], 'final', attempt=1)
    assert result['status'] == 'needs_attention'
    assert 'missing' in result['message'].lower()
    assert result['artifacts'] == []


def test_valid_output_is_opened_and_retained_not_fact_verified(tmp_path):
    svc = service(tmp_path); task = draft(svc)
    asyncio.run(svc.start(task['id'], task['version'], 'one', lambda *_: None))
    output = tmp_path / task['output_path']; output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text('# Observations\n\nThe supplied notes describe a small pilot.\n')
    result = svc.event(task['id'], 'final', attempt=1)
    assert result['status'] == 'ready_for_review'
    artifact = result['artifacts'][0]
    assert 'factual' in artifact['verification'].lower()
    output.write_text('changed')
    assert svc.artifact(task['id'], artifact['id'])[1].startswith(b'# Observations')


def test_malformed_data_and_unsafe_sources(tmp_path):
    svc = service(tmp_path)
    for path in ['../outside.txt', '/etc/passwd', '.dream/workflows/tasks.sqlite3']:
        with pytest.raises(ValueError):
            draft(svc, sources=path)
    (tmp_path / 'sample.csv').write_text('name,value\na,1\n')
    task = svc.create('analysis', {'goal': 'Analyze sample', 'sources': 'sample.csv'})
    asyncio.run(svc.start(task['id'], task['version'], 'one', lambda *_: None))
    p = tmp_path / task['output_path']; p.parent.mkdir(parents=True, exist_ok=True); p.write_text('{broken')
    assert svc.event(task['id'], 'final', attempt=1)['status'] == 'needs_attention'


def test_uncertain_dispatch_and_explicit_recovery_never_replay(tmp_path):
    svc = service(tmp_path); task = draft(svc); calls=[]
    def fail(*_):
        calls.append(1); raise RuntimeError('connection lost after enqueue')
    state = asyncio.run(svc.start(task['id'], task['version'], 'one', fail))
    assert state['status'] == 'unknown'
    state = svc.recover(task['id'], state['version'])
    assert state['status'] == 'interrupted' and len(calls) == 1
    retry = svc.revise(task['id'], state['version'])
    assert retry['attempt'] == 2 and retry['status'] == 'draft'
    assert retry['output_path'] != task['output_path']
    assert svc.event(task['id'], 'final', attempt=1)['status'] == 'draft'


def test_symlink_output_rejected(tmp_path):
    svc=service(tmp_path); task=draft(svc)
    asyncio.run(svc.start(task['id'], task['version'], 'one', lambda *_: None))
    p=tmp_path/task['output_path'];p.parent.mkdir(parents=True,exist_ok=True)
    target=tmp_path/'source.md';target.write_text('# Should not follow\nSome content here.')
    p.symlink_to(target)
    assert svc.event(task['id'], 'final', attempt=1)['status']=='needs_attention'


def test_concurrent_duplicate_starts_dispatch_once(tmp_path):
    svc=service(tmp_path);task=draft(svc)
    async def scenario():
        entered=asyncio.Event();release=asyncio.Event();calls=[]
        async def dispatch(*args):
            calls.append(args);entered.set();await release.wait()
        first=asyncio.create_task(svc.start(task['id'],1,'same-request',dispatch))
        await entered.wait()
        duplicate=await svc.start(task['id'],1,'same-request',dispatch)
        assert duplicate['status']=='queued'
        release.set();await first
        assert len(calls)==1
    asyncio.run(scenario())


def test_valid_analysis_and_missing_data_input(tmp_path):
    svc=service(tmp_path)
    with pytest.raises(ValueError,match='requires'):
        svc.create('analysis',{'goal':'Find average'})
    (tmp_path/'input.csv').write_text('value\n2\n4\n')
    task=svc.create('analysis',{'goal':'Find average','sources':'input.csv'})
    asyncio.run(svc.start(task['id'],1,'one',lambda *_:None))
    p=tmp_path/task['output_path'];p.parent.mkdir(parents=True)
    p.write_text(json.dumps({'summary':'Average is 3','results':[{'mean':3}],'sources':['input.csv']}))
    result=svc.event(task['id'],'final',attempt=1)
    assert result['status']=='ready_for_review'
    assert 'calculations' in result['message']


def test_pptx_rejects_fake_package_and_accepts_real_slides(tmp_path):
    import io
    import zipfile
    from dream.workflows.validation import verify
    fake=io.BytesIO()
    with zipfile.ZipFile(fake,'w') as z:
        for path in ['[Content_Types].xml','_rels/.rels','ppt/presentation.xml','ppt/_rels/presentation.xml.rels']:
            z.writestr(path,'<fake/>')
        z.writestr('ppt/slides/slide1.xml','<fake xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"><a:t>fake</a:t></fake>')
    with pytest.raises(ValueError):
        verify(fake.getvalue(),'presentation')
    from pptx import Presentation
    presentation=Presentation();slide=presentation.slides.add_slide(presentation.slide_layouts[0]);slide.shapes.title.text='Source observations'
    output=io.BytesIO();presentation.save(output)
    assert 'slide' in verify(output.getvalue(),'presentation')


def test_draft_and_recovery_survive_new_service(tmp_path):
    svc=service(tmp_path);task=draft(svc)
    asyncio.run(svc.start(task['id'],1,'one',lambda *_:None))
    restored=service(tmp_path)
    assert restored.get(task['id'])['status']=='queued'
    state=restored.get(task['id'])
    assert restored.recover(task['id'],state['version'])['status']=='interrupted'


def test_oversized_and_parent_symlink_sources_rejected(tmp_path):
    svc=service(tmp_path)
    large=tmp_path/'large.txt'
    with large.open('wb') as f:f.truncate(8*1024*1024+1)
    with pytest.raises(ValueError,match='8 MiB'):draft(svc,sources='large.txt')
    folder=tmp_path/'real';folder.mkdir();(folder/'source.txt').write_text('source')
    (tmp_path/'alias').symlink_to(folder,target_is_directory=True)
    with pytest.raises(ValueError):draft(svc,sources='alias/source.txt')


def test_cancelled_dispatch_keeps_uncertainty_visible(tmp_path):
    svc=service(tmp_path);task=draft(svc)
    async def dispatch(*_):raise asyncio.CancelledError()
    with pytest.raises(asyncio.CancelledError):
        asyncio.run(svc.start(task['id'],1,'cancelled',dispatch))
    assert svc.get(task['id'])['status']=='unknown'


def test_checker_rejects_absent_content_and_private_state_symlink(tmp_path):
    from dream.workflows.validation import verify
    for data,kind in [(b'# Header','report'),(b'{"summary":"a","results":[],"sources":[]}', 'analysis'),(b'not a zip','presentation')]:
        with pytest.raises(ValueError):verify(data,kind)
    outside=tmp_path/'outside';outside.mkdir();(tmp_path/'.dream').symlink_to(outside,target_is_directory=True)
    with pytest.raises(ValueError,match='symlinks'):service(tmp_path)


def test_report_heading_alone_and_empty_result_records_fail_content_check():
    from dream.workflows.validation import verify
    with pytest.raises(ValueError):verify(b'# ' + b'Heading ' * 10,'report')
    with pytest.raises(ValueError):verify(b'{"summary":"Claim","results":[{}],"sources":["input.csv"]}','analysis')


def test_store_rejects_database_symlink_swapped_after_initialization(tmp_path):
    import sqlite3
    svc=service(tmp_path);task=draft(svc)
    outside=tmp_path/'outside.sqlite'
    with sqlite3.connect(outside) as db:db.execute('CREATE TABLE sentinel(value TEXT)')
    svc.store.path.rename(svc.store.path.with_suffix('.original'))
    svc.store.path.symlink_to(outside)
    with pytest.raises(ValueError):svc.list()
    with sqlite3.connect(outside) as db:
        assert db.execute("SELECT name FROM sqlite_master WHERE type='table'").fetchall()==[('sentinel',)]


def test_private_state_permissions(tmp_path):
    import stat
    svc=service(tmp_path);draft(svc)
    assert stat.S_IMODE((tmp_path/'.dream').stat().st_mode)==0o700
    assert stat.S_IMODE(svc.store.path.parent.stat().st_mode)==0o700
    assert stat.S_IMODE(svc.store.path.stat().st_mode)==0o600


@pytest.mark.parametrize('value',['NaN','Infinity','-Infinity','1e999'])
def test_analysis_rejects_nonfinite_json_numbers(value):
    from dream.workflows.validation import verify
    data=('{"summary":"Claim","results":[{"nested":[{"value":'+value+'}]}],"sources":["input.csv"]}').encode()
    with pytest.raises(ValueError):verify(data,'analysis')


def test_pptx_requires_correct_main_content_type():
    import io,zipfile
    from pptx import Presentation
    from dream.workflows.validation import verify
    deck=Presentation();slide=deck.slides.add_slide(deck.slide_layouts[0]);slide.shapes.title.text='Fixture'
    original=io.BytesIO();deck.save(original)
    malformed=io.BytesIO()
    with zipfile.ZipFile(original) as source,zipfile.ZipFile(malformed,'w') as target:
        for entry in source.infolist():
            data=source.read(entry)
            if entry.filename=='[Content_Types].xml':
                data=data.replace(b'application/vnd.openxmlformats-officedocument.presentationml.presentation.main+xml',b'application/invalid')
            target.writestr(entry,data)
    with pytest.raises(ValueError):verify(malformed.getvalue(),'presentation')


def test_store_detects_symlink_swap_during_sqlite_connect(tmp_path,monkeypatch):
    import sqlite3
    svc=service(tmp_path);draft(svc)
    outside=tmp_path/'external.sqlite'
    with sqlite3.connect(outside) as db:db.execute('CREATE TABLE sentinel(value TEXT)')
    original=sqlite3.connect;before=outside.read_bytes()
    def swapped(*args,**kwargs):
        svc.store.path.rename(svc.store.path.with_suffix('.held'))
        svc.store.path.symlink_to(outside)
        return original(*args,**kwargs)
    monkeypatch.setattr(sqlite3,'connect',swapped)
    with pytest.raises(ValueError,match='changed'):svc.list()
    assert outside.read_bytes()==before


def test_store_and_source_reject_replaced_parent_directories(tmp_path):
    svc=service(tmp_path);draft(svc)
    outside=tmp_path/'outside';outside.mkdir()
    root=tmp_path/'.dream'/'workflows';root.rename(root.with_name('held'))
    root.symlink_to(outside,target_is_directory=True)
    with pytest.raises(ValueError):svc.list()
    assert not list(outside.iterdir())
